"""서비스 레이어 (core/services.py)

app.py(UI)는 이 서비스들만 호출한다. 비즈니스 로직은 전부 여기(와 core/*)에.
향후 REST API를 만들 때도 같은 서비스를 재사용하면 된다.
"""

from __future__ import annotations

import io
import json
import random

import pandas as pd

from config import (DEFAULT_BATCH_SIZE, METHODOLOGY_NOTE, SAMPLE_Q_PATH,
                    STRATEGY_ORDER)
from .persona import sample_persona
from .question_types import get_question_type
from .strategies import (GenerationContext, get_strategy, select_strategy)


# ----------------------------------------------------------------------------
# 설문 문항 관리
# ----------------------------------------------------------------------------
class QuestionnaireService:
    @staticmethod
    def new_question_id(questions: list[dict]) -> str:
        nums = [int(q["id"][1:]) for q in questions
                if q["id"].startswith("Q") and q["id"][1:].isdigit()]
        return f"Q{(max(nums) + 1) if nums else 1}"

    @staticmethod
    def question_summary(q: dict) -> str:
        qtype = get_question_type(q["type"])
        base = f"{q['id']}. {q['text']} ({qtype.label})"
        if qtype.needs_options and q.get("options"):
            base += " — " + ", ".join(qtype.options_for(q))
        if q.get("branching"):
            base += " ⤷분기"
        return base

    @staticmethod
    def validate_question(q: dict) -> list[str]:
        """문항 유효성 검사. 오류 메시지 목록 (빈 리스트 = 유효)."""
        return get_question_type(q["type"]).validate(q)

    @staticmethod
    def load_sample_questions() -> list[dict]:
        return json.loads(SAMPLE_Q_PATH.read_text(encoding="utf-8"))


# ----------------------------------------------------------------------------
# 분포 학습
# ----------------------------------------------------------------------------
class LearningService:
    @staticmethod
    def _match_column(df: pd.DataFrame, q: dict) -> str | None:
        """엑셀 컬럼을 문항에 매칭: 문항 텍스트 일치 → Q번호 일치 순."""
        cols = list(df.columns)
        if q["text"] in cols:
            return q["text"]
        if q["id"] in cols:
            return q["id"]
        qid_lower = q["id"].lower()
        for c in cols:
            if str(c).strip().lower() == qid_lower:
                return c
        return None

    def learn_distributions(self, df: pd.DataFrame,
                            questions: list[dict]) -> dict:
        """실제 응답 엑셀에서 문항별 응답 분포를 추출한다.

        Returns: {qid: {"matched_col": str, "type": str,
                         "dist": {option: prob}, "n": int}}
        주관식은 분포 없이 n만 기록한다.
        """
        learned: dict = {}
        for q in questions:
            qtype = get_question_type(q["type"])
            col = self._match_column(df, q)
            if col is None:
                continue
            series = df[col].dropna()
            n = len(series)
            if n == 0:
                continue
            dist = qtype.learn_distribution(q, series)
            if dist is None:
                continue
            learned[q["id"]] = {
                "matched_col": str(col),
                "type": q["type"],
                "dist": {str(k): float(v) for k, v in dist.items()},
                "n": n,
            }
        return learned


# ----------------------------------------------------------------------------
# 가상 응답 생성 (전략 오케스트레이션)
# ----------------------------------------------------------------------------
class GenerationService:
    def __init__(self, strategy_order: list[str] | None = None) -> None:
        self.strategy_order = strategy_order or STRATEGY_ORDER

    def _select_strategy(self, q: dict, ctx: GenerationContext):
        return select_strategy(q, ctx, self.strategy_order)

    def generate_responses(
        self,
        questions: list[dict],
        persona_dims: dict[str, dict[str, float]],
        n: int,
        *,
        provider_id: str,
        model: str,
        learned: dict | None = None,
        batch_size: int = DEFAULT_BATCH_SIZE,
        progress_cb=None,
        cache: dict | None = None,
        seed: int | None = None,
        transcripts: list[dict] | None = None,
        use_precision: bool = False,
        generate_fn=None,
    ) -> pd.DataFrame:
        """가상 응답 n건을 생성해 DataFrame으로 반환.

        문항마다 레지스트리 우선순위대로 전략을 매칭한다:
        - 학습된 분포가 있는 선택형 → 분포학습 (API 호출 없음)
        - 고정밀 컨텍스트의 선택형 → 고정밀 (인터뷰 기반)
        - 나머지 선택형 → 분포추출 (Prolific 2026 방식)
        - 주관식 → 주관식 생성 (LLM 배치)
        """
        rng = random.Random(seed)
        learned = learned or {}
        transcripts = transcripts or []
        ctx = GenerationContext(
            provider_id=provider_id,
            model=model,
            learned=learned,
            transcripts=transcripts,
            use_precision=bool(use_precision and transcripts),
            batch_size=batch_size,
            cache=cache if cache is not None else {},
            generate_fn=generate_fn,
        )

        # 문항 → 전략 매칭 (전략별 그룹)
        groups: dict[str, list[dict]] = {}
        for q in questions:
            strategy = self._select_strategy(q, ctx)
            groups.setdefault(strategy.name, []).append(q)

        rows: list[dict] = []
        total_batches = (n + batch_size - 1) // batch_size
        for b in range(total_batches):
            bsize = min(batch_size, n - b * batch_size)
            personas = [sample_persona(persona_dims, rng) for _ in range(bsize)]

            sampled: list[dict] = [{} for _ in range(bsize)]
            for sname, qs in groups.items():
                strategy = get_strategy(sname)
                results = strategy.generate(ctx, qs, personas, rng)
                for i, ans in enumerate(results):
                    sampled[i].update(ans)

            for i, persona in enumerate(personas):
                row = {f"프로필_{dim}": val for dim, val in persona.items()}
                for q in questions:
                    row[q["id"]] = sampled[i].get(q["id"], "")
                rows.append(row)

            if progress_cb:
                progress_cb((b + 1) / total_batches)

        col_order = [f"프로필_{dim}" for dim in persona_dims] + [q["id"] for q in questions]
        df = pd.DataFrame(rows)
        return df[[c for c in col_order if c in df.columns]]


# ----------------------------------------------------------------------------
# 엑셀 출력
# ----------------------------------------------------------------------------
class ExportService:
    @staticmethod
    def to_excel_bytes(df: pd.DataFrame, sheet_name: str = "가상응답",
                       meta: dict | None = None,
                       diagnostics: list[dict] | None = None) -> bytes:
        from core.diagnostics import findings_to_dataframe
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            info_rows = [["항목", "내용"]]
            for line in METHODOLOGY_NOTE.strip().split("\n"):
                info_rows.append([line])
            if meta:
                info_rows.append([""])
                for k, v in meta.items():
                    info_rows.append([k, str(v)])
            pd.DataFrame(info_rows).to_excel(
                writer, index=False, header=False, sheet_name="방법론_안내"
            )
            # 파일럿 진단 리포트 시트 (규칙 기반, 환각 없음)
            findings_to_dataframe(diagnostics or []).to_excel(
                writer, index=False, sheet_name="진단_리포트"
            )
        return buf.getvalue()

    @staticmethod
    def to_word_bytes(df: pd.DataFrame,
                      meta: dict | None = None,
                      diagnostics: list[dict] | None = None,
                      questions: list[dict] | None = None) -> bytes:
        """파일럿 진단 리포트를 Word(.docx)로 생성.

        깔끔한 디자인: 제목, 메타 정보, 진단 결과 표, 방법론 안내 섹션.
        """
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from datetime import datetime

        doc = Document()

        # 기본 폰트 설정
        style = doc.styles["Normal"]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(10)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

        # 제목
        title = doc.add_heading("파일럿 테스트 진단 리포트", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 생성 일시
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        doc.add_paragraph()  # 여백

        # 메타 정보
        if meta:
            doc.add_heading("기본 정보", level=1)
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            hdr[0].text = "항목"
            hdr[1].text = "내용"
            for k, v in meta.items():
                row = table.add_row().cells
                row[0].text = str(k)
                row[1].text = str(v)
            doc.add_paragraph()

        # 진단 결과
        doc.add_heading("진단 결과", level=1)
        findings = diagnostics or []
        if not findings:
            p = doc.add_paragraph()
            run = p.add_run("✅ 진단 결과: 문제가 발견되지 않았습니다. 설문지가 정상적으로 작동합니다.")
            run.font.color.rgb = RGBColor(0x1B, 0x7A, 0x34)
        else:
            # 요약
            high = sum(1 for f in findings if f.get("심각도") == "높음")
            mid = len(findings) - high
            p = doc.add_paragraph()
            p.add_run(f"총 {len(findings)}건 발견 (높음 {high}건, 중간 {mid}건)")

            table = doc.add_table(rows=1, cols=5)
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["문항", "문제", "근거", "해결책", "심각도"]
            for i, h in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell.text = h
                for par in cell.paragraphs:
                    for r in par.runs:
                        r.bold = True

            for f in findings:
                row = table.add_row().cells
                row[0].text = str(f.get("문항", ""))
                row[1].text = str(f.get("문제", ""))
                row[2].text = str(f.get("근거", ""))
                row[3].text = str(f.get("해결책", ""))
                sev = str(f.get("심각도", ""))
                row[4].text = sev
                # 심각도에 따라 색상
                for par in row[4].paragraphs:
                    for r in par.runs:
                        r.bold = True
                        if sev == "높음":
                            r.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
                        else:
                            r.font.color.rgb = RGBColor(0xCC, 0x7A, 0x00)

            # 열 너비 조정
            for row in table.rows:
                row.cells[0].width = Cm(3)
                row.cells[1].width = Cm(4)
                row.cells[2].width = Cm(4)
                row.cells[3].width = Cm(4)
                row.cells[4].width = Cm(2)

        doc.add_paragraph()

        # 응답 요약 통계
        if df is not None and not df.empty and questions:
            doc.add_heading("문항별 응답 요약", level=1)
            for q in questions:
                qid = q.get("id")
                if qid not in df.columns:
                    continue
                doc.add_heading(q.get("text", qid), level=2)
                vc = df[qid].astype(str).value_counts()
                table = doc.add_table(rows=1, cols=3)
                table.style = "Light Grid Accent 1"
                hdr = table.rows[0].cells
                hdr[0].text = "응답"
                hdr[1].text = "건수"
                hdr[2].text = "비율"
                for val, cnt in vc.items():
                    row = table.add_row().cells
                    row[0].text = str(val)
                    row[1].text = str(cnt)
                    row[2].text = f"{cnt/len(df)*100:.1f}%"
                doc.add_paragraph()

        # 방법론 안내
        doc.add_heading("방법론 안내", level=1)
        for line in METHODOLOGY_NOTE.strip().split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip(), style="List Bullet")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()


# ----------------------------------------------------------------------------
# PDF 설문지 자동 인식
# ----------------------------------------------------------------------------
class PdfQuestionnaireService:
    """PDF 설문지를 업로드하면 AI가 문항을 자동 파싱합니다.

    분기문항(스킵 로직)도 인식하여 branching 필드에 저장합니다.
    """

    @staticmethod
    def extract_text(pdf_bytes: bytes) -> str:
        """PDF에서 텍스트를 추출합니다."""
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(pdf_bytes))
        parts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
        return "\n\n".join(parts)

    @staticmethod
    def parse_questions(pdf_text: str, provider_id: str = "gemini") -> list[dict]:
        """AI를 이용해 PDF 텍스트에서 설문 문항을 파싱합니다.

        Returns:
            [{"text": ..., "type": "single|multi|likert|open",
              "options": [...], "branching": "..."}, ...]
            branching은 분기 로직 설명 문자열 (없으면 "")
        """
        from .providers import generate

        prompt = f"""다음은 설문지 PDF에서 추출한 텍스트입니다. 설문 문항을 JSON 배열로 파싱하세요.

규칙:
1. 각 문항은 {{"text": "문항 내용", "type": "single|multi|likert|open", "options": ["보기1", "보기2", ...], "branching": "분기 설명 또는 빈 문자열"}} 형태입니다.
2. type 판단 기준:
   - single: 보기 중 하나만 선택 (단일선택)
   - multi: 복수 선택 가능 (다중선택, "모두 선택" 등)
   - likert: 5점 척도 (매우 그렇다~전혀 그렇지 않다 등)
   - open: 주관식 서술형
3. 분기문항(스킵 로직)이 있으면 branching에 설명을 적으세요. 예: "문3에서 ① 응답 시 문5로 이동", "문2의 ② 응답자만 응답"
   분기가 없으면 빈 문자열("")을 넣으세요.
4. 보기 번호(①②③, 1. 2. 3. 등)는 제거하고 보기 내용만 추출하세요.
5. 설문 안내문, 표지, 인구통계 안내 등은 문항에서 제외하세요.
6. JSON 배열만 출력하고 다른 설명은 붙이지 마세요.

--- 설문지 텍스트 ---
{pdf_text[:12000]}
--- 끝 ---"""

        result = generate(provider_id, prompt)
        # JSON 추출 (코드블록 제거)
        result = result.strip()
        if result.startswith("```"):
            lines = result.split("\n")
            result = "\n".join(lines[1:-1] if lines[-1].strip() == "```" else lines[1:])
        try:
            questions = json.loads(result)
        except json.JSONDecodeError:
            # JSON 파싱 실패 시 빈 리스트
            return []
        # 유효성 검증 및 정규화
        valid_types = {"single", "multi", "likert", "open"}
        cleaned = []
        for q in questions:
            if not isinstance(q, dict) or not q.get("text", "").strip():
                continue
            qtype = q.get("type", "single")
            if qtype not in valid_types:
                qtype = "single"
            options = q.get("options", [])
            if not isinstance(options, list):
                options = []
            options = [str(o).strip() for o in options if str(o).strip()]
            cleaned.append({
                "text": q["text"].strip(),
                "type": qtype,
                "options": options,
                "branching": str(q.get("branching", "")).strip(),
            })
        return cleaned
