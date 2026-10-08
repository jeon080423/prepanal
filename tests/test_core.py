"""core 로직 테스트 (실제 LLM 호출 없음 — generate_fn 주입)."""
import io
import json
import random
import re

import pandas as pd
import pytest

from config import (DISCLAIMER, METHODOLOGY_NOTE, SAMPLE_Q_PATH,
                    SAMPLE_XLSX_PATH)
from core import (ExportService, GenerationService, LearningService,
                  MemoryStore, QuestionnaireService, get_question_type,
                  sample_persona)
from core.strategies import GenerationContext, parse_json_array


@pytest.fixture
def sample_questions():
    return json.loads(SAMPLE_Q_PATH.read_text(encoding="utf-8"))


@pytest.fixture
def persona_dims():
    return {"성별": {"남성": 50.0, "여성": 50.0},
            "연령대": {"20대": 50.0, "30대": 50.0}}


# ----------------------------------------------------------------------------
# 페르소나
# ----------------------------------------------------------------------------
def test_persona_sampling_deterministic(persona_dims):
    a = [sample_persona(persona_dims, random.Random(7)) for _ in range(20)]
    b = [sample_persona(persona_dims, random.Random(7)) for _ in range(20)]
    assert a == b


def test_persona_respects_distribution():
    dims = {"성별": {"남성": 90.0, "여성": 10.0}}
    rng = random.Random(1)
    vals = [sample_persona(dims, rng)["성별"] for _ in range(500)]
    assert vals.count("남성") / 500 > 0.8


# ----------------------------------------------------------------------------
# 분포 학습 (샘플 엑셀)
# ----------------------------------------------------------------------------
def test_learn_distributions_sample_excel(sample_questions):
    df = pd.read_excel(SAMPLE_XLSX_PATH)
    learned = LearningService().learn_distributions(df, sample_questions)
    assert len(learned) == 6  # 6문항 전수 매칭
    for qid, info in learned.items():
        assert info["n"] > 0
        if info["dist"]:
            assert abs(sum(info["dist"].values()) - 1.0) < 1e-6
    assert learned["Q5"]["dist"] == {}  # 주관식은 분포 없음


def test_learn_distributions_no_match():
    df = pd.DataFrame({"전혀다른컬럼": ["a", "b"]})
    qs = [{"id": "Q1", "text": "매칭안됨", "type": "single",
           "options": ["a", "b"]}]
    assert LearningService().learn_distributions(df, qs) == {}


# ----------------------------------------------------------------------------
# 문항 서비스
# ----------------------------------------------------------------------------
def test_new_question_id():
    assert QuestionnaireService.new_question_id([]) == "Q1"
    qs = [{"id": "Q1"}, {"id": "Q3"}]
    assert QuestionnaireService.new_question_id(qs) == "Q4"


def test_question_summary():
    q = {"id": "Q1", "text": "어디?", "type": "single",
         "options": ["A", "B"]}
    s = QuestionnaireService.question_summary(q)
    assert "Q1" in s and "단일선택" in s and "A, B" in s


# ----------------------------------------------------------------------------
# JSON 파싱
# ----------------------------------------------------------------------------
def test_parse_json_array_fences():
    raw = '```json\n[{"Q1": "A"}, {"Q1": "B"}]\n```'
    assert parse_json_array(raw) == [{"Q1": "A"}, {"Q1": "B"}]
    assert parse_json_array('[{"Q1": "A"}]') == [{"Q1": "A"}]
    with pytest.raises(ValueError):
        parse_json_array("응답 없음")


# ----------------------------------------------------------------------------
# 생성 파이프라인 (가짜 LLM)
# ----------------------------------------------------------------------------
def _profile_count(prompt: str) -> int:
    return len(re.findall(r"(?m)^\d+\.\s", prompt))


def fake_generate(provider_id, prompt, model):
    n = _profile_count(prompt)
    assert n > 0
    if "[주관식 문항]" in prompt:
        return json.dumps([{"Q3": "배송이 빨라서 좋아요"} for _ in range(n)],
                          ensure_ascii=False)
    # 분포 추출 프롬프트
    return json.dumps([{"Q2": {"예": 0.7, "아니오": 0.3}} for _ in range(n)],
                      ensure_ascii=False)


@pytest.fixture
def mixed_questions():
    return [
        {"id": "Q1", "text": "학습된 문항", "type": "single",
         "options": ["A", "B"]},
        {"id": "Q2", "text": "미학습 문항", "type": "single",
         "options": ["예", "아니오"]},
        {"id": "Q3", "text": "주관식 문항", "type": "open", "options": []},
    ]


def test_generate_mixed_strategies(mixed_questions, persona_dims):
    learned = {"Q1": {"matched_col": "c", "type": "single",
                      "dist": {"A": 1.0, "B": 0.0}, "n": 50}}
    svc = GenerationService()
    df = svc.generate_responses(
        mixed_questions, persona_dims, 4,
        provider_id="gemini", model="dummy",
        learned=learned, batch_size=4, seed=42,
        generate_fn=fake_generate,
    )
    assert list(df.columns) == ["프로필_성별", "프로필_연령대",
                                "Q1", "Q2", "Q3"]
    assert len(df) == 4
    assert set(df["Q1"].unique()) <= {"A", "B"}
    assert set(df["Q2"].unique()) <= {"예", "아니오"}
    assert (df["Q3"] == "배송이 빨라서 좋아요").all()


def test_generate_precision_mode(mixed_questions, persona_dims):
    transcripts = [{"name": "t1", "text": "인터뷰 내용 " * 100}]
    svc = GenerationService()
    df = svc.generate_responses(
        mixed_questions, persona_dims, 3,
        provider_id="gemini", model="dummy",
        batch_size=3, seed=1,
        transcripts=transcripts, use_precision=True,
        generate_fn=fake_generate,
    )
    assert len(df) == 3
    assert (df["Q3"] == "배송이 빨라서 좋아요").all()


def test_generate_no_llm_when_all_learned(mixed_questions, persona_dims):
    learned = {
        "Q1": {"matched_col": "c", "type": "single",
               "dist": {"A": 0.5, "B": 0.5}, "n": 10},
        "Q2": {"matched_col": "c", "type": "single",
               "dist": {"예": 0.5, "아니오": 0.5}, "n": 10},
    }
    calls = []

    def counting_fake(pid, prompt, model):
        calls.append(prompt)
        return fake_generate(pid, prompt, model)

    df = GenerationService().generate_responses(
        mixed_questions, persona_dims, 4,
        provider_id="gemini", model="dummy",
        learned=learned, batch_size=4, seed=3,
        generate_fn=counting_fake,
    )
    assert len(df) == 4
    # Q1·Q2는 학습 분포 → LLM 호출은 주관식(Q3) 1건만
    assert len(calls) == 1
    assert "[주관식 문항]" in calls[0]


# ----------------------------------------------------------------------------
# 엑셀 출력
# ----------------------------------------------------------------------------
def test_excel_export_contains_methodology_sheet():
    df = pd.DataFrame({"Q1": ["A", "B"], "Q2": ["예", "아니오"]})
    data = ExportService.to_excel_bytes(df, meta={"AI": "테스트"})
    assert data[:2] == b"PK"  # zip/xlsx 매직바이트
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert "가상응답" in wb.sheetnames
    assert "방법론_안내" in wb.sheetnames
    ws = wb["방법론_안내"]
    texts = " ".join(str(c.value) for row in ws.iter_rows() for c in row
                     if c.value)
    assert "파인튜닝" in texts
    assert DISCLAIMER[:20] in texts
    assert "테스트" in texts  # meta 포함


# ----------------------------------------------------------------------------
# Store
# ----------------------------------------------------------------------------
def test_memory_store():
    s = MemoryStore()
    assert s.get("missing", "dflt") == "dflt"
    assert "k" not in s
    assert s.set_default("k", 1) == 1
    assert s.set_default("k", 2) == 1  # 기존값 유지
    s.delete("k")
    assert "k" not in s


def test_methodology_note_mentions_research():
    assert "Prolific 2026" in METHODOLOGY_NOTE
    assert "83%" in METHODOLOGY_NOTE
